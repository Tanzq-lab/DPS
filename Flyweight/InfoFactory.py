import os
import random
from typing import Dict
import openpyxl
from docx import Document

from Flyweight.PersonalDetail import PersonalDetail


class InfoFactory:
    """
    该类主要的作用就是返回对应的信息,并创建指定的类型的文件
    """
    _flock: Dict[str, PersonalDetail] = {}

    # 对工厂进行初始化
    def __init__(self):
        print("正在加载文件数据...")
        for i, j, k in os.walk('Data'):
            for fileName in k:
                if fileName.split('.')[-1] != 'bak' and fileName.split('.')[-1] != 'dir':
                    self._flock[fileName.split('_')[0]] = PersonalDetail("Data\\" + fileName)

        print("数据加载完成！")

    # name => 用户名   extname => 后缀名
    def Builder(self, name: str, extname: str):
        # 如果之前没有出现过该数据的化,就重新生成一个数据,并加入到字典中.
        if not self._flock.get(name):
            print("之前没有出现过该数据,现已新建~")
            self._flock[name] = PersonalDetail()
            self._flock[name].name = name

        # 如果之前保存过该数据,如果覆盖就将目前的数据进行清空, 修改不用处理,然后将对应的序列化数据删除.
        # 因为待会重新生成这个数据,所以说这个数据就相当于没有用了,直接删除.
        else:
            PD = self._flock[name]
            opt = input("该数据之前出现过,请问您如何处理? (C)覆盖,(M)修改 : ")
            if opt == "C":
                PD.clear()

            print(PD.Type)
            if PD.Type == "Tanzq":
                fileFullPath = "Data\\" + name + '_shelve.Tanzq'
                os.remove(fileFullPath + ".bak")
                os.remove(fileFullPath + ".dat")
                os.remove(fileFullPath + ".dir")

            elif PD.Type == 'txt':
                fileFullPath = "Data\\" + name + '_pickle.txt'
                os.remove(fileFullPath)

            elif PD.Type == 'json':
                fileFullPath = "Data\\" + name + '_json.json'
                os.remove(fileFullPath)

        # 创建指定的文件.
        return self._flock[name].createFile(extname)

    def serialization(self, fileFullPath: str):
        fileType = fileFullPath.split('.')[-1]
        PD = self._flock[fileFullPath.split('.')[0].split('\\')[-1]]

        # print(fileType)
        if fileType == 'xlsx':
            # 获取到表格对象
            workbook = openpyxl.load_workbook(fileFullPath)
            worksheet = workbook.active

            # 得到对应的值
            PD.name = list(worksheet.rows)[1][0].value
            PD.address = list(worksheet.rows)[1][1].value
            PD.phone = list(worksheet.rows)[1][2].value
            # print(PD.__dict__)

        elif fileType == 'txt':
            with open(fileFullPath, "r") as fp:
                # 得到对应的值
                info = fp.readlines()
                PD.name = info[0].replace('姓名：', '').replace('\n', '')
                PD.address = info[1].replace('地址：', '').replace('\n', '')
                PD.phone = info[2].replace('电话：', '')
                # print(PD.__dict__)

        elif fileType == 'docx':
            doc = Document(fileFullPath)
            table = doc.tables[0]
            PD.name = table.rows[0].cells[1].text
            PD.address = table.rows[1].cells[1].text
            PD.phone = table.rows[2].cells[1].text

        # 序列化的方式使用随机化的形式
        typeList = ["json", "txt", "Tanzq"]
        PD.Type = typeList[random.randint(0, 2)]

        # 数据已经完成, 最后进行序列化处理
        PD.encoding()

        os.remove(fileFullPath)
